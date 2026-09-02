"""
services/reports.py
Excel export generation — bilingual (EN/ES).
All public functions accept an optional lang='en'|'es' parameter.
"""
import io
import calendar
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from db import get_conn

# ── Translations for report headers ──────────────────────────────────────────
_T = {
    "en": {
        "employee_id":    "Employee ID",
        "name":           "Name",
        "department":     "Department",
        "days_present":   "Days Present",
        "attendance_pct": "Attendance %",
        "in_count":       "IN Count",
        "out_count":      "OUT Count",
        "days":           "Days",
        "total_hours":    "Total Hours",
        "in":             "IN",
        "out":            "OUT",
        "weekly_title":   "Weekly Attendance",
        "monthly_title":  "Monthly Summary",
        "no_dept":        "No Department",
        "to":             "to",
    },
    "es": {
        "employee_id":    "ID Empleado",
        "name":           "Nombre",
        "department":     "Departamento",
        "days_present":   "Días Presentes",
        "attendance_pct": "% Asistencia",
        "in_count":       "Entradas",
        "out_count":      "Salidas",
        "days":           "Días",
        "total_hours":    "Total Horas",
        "in":             "ENTRADA",
        "out":            "SALIDA",
        "weekly_title":   "Asistencia",
        "monthly_title":  "Resumen Mensual",
        "no_dept":        "Sin Departamento",
        "to":             "al",
    },
}

_MONTHS_ES = ["Enero","Febrero","Marzo","Abril","Mayo","Junio",
              "Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
_DAYS_ES   = ["Lun","Mar","Mié","Jue","Vie","Sáb","Dom"]
_DAYS_EN   = ["Mon","Tue","Wed","Thu","Fri","Sat","Sun"]


def _tr(lang: str, key: str) -> str:
    return _T.get(lang, _T["en"]).get(key, key)


def _day_label(d: date, lang: str) -> str:
    if lang == "es":
        return f"{_DAYS_ES[d.weekday()]} {d.day:02d}/{d.month:02d}"
    return d.strftime("%a %d/%m")


def _month_label(year: int, month: int, lang: str) -> str:
    if lang == "es":
        return f"{_MONTHS_ES[month-1]} {year}"
    from datetime import date as _date
    return _date(year, month, 1).strftime("%B %Y")


def _thin_border():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)

def _header_fill():
    return PatternFill("solid", fgColor="1E293B")

def _alt_fill():
    return PatternFill("solid", fgColor="F8FAFC")

def _apply_header(ws, row, cols):
    for i, label in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = _header_fill()
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _thin_border()
    ws.row_dimensions[row].height = 20


# ── Weekly FIFO export ────────────────────────────────────────────────────────

def export_weekly_fifo(start_date: date, end_date: date,
                       lang: str = "en",
                       emp_ids: list = None) -> io.BytesIO:
    """First-IN / Last-OUT grid export. Supports quincena period too."""
    from services.attendance import get_weekly_attendance

    week_days = []
    d = start_date
    while d <= end_date:
        week_days.append(d)
        d += timedelta(days=1)

    data = get_weekly_attendance(start_date, end_date)

    # Apply employee filter (department filter)
    if emp_ids:
        data = {k: v for k, v in data.items() if str(k) in emp_ids}

    wb = Workbook()
    ws = wb.active
    ws.title = _tr(lang, "weekly_title")[:31]

    # Title row
    total_cols = 2 + len(week_days) * 2 + 1
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=total_cols)
    title = (f"{_tr(lang, 'weekly_title')} — "
             f"{start_date.strftime('%d/%m/%Y')} "
             f"{_tr(lang, 'to')} {end_date.strftime('%d/%m/%Y')}")
    c = ws.cell(1, 1, title)
    c.font = Font(bold=True, size=12, color="FFFFFF")
    c.fill = _header_fill()
    c.alignment = Alignment(horizontal="center")

    # Row 2 — ID, Name, then day columns
    ws.cell(2, 1, _tr(lang, "employee_id")).font = Font(bold=True, color="FFFFFF")
    ws.cell(2, 1).fill = _header_fill()
    ws.cell(2, 2, _tr(lang, "name")).font = Font(bold=True, color="FFFFFF")
    ws.cell(2, 2).fill = _header_fill()
    col = 3
    for d in week_days:
        ws.merge_cells(start_row=2, start_column=col,
                       end_row=2, end_column=col+1)
        c = ws.cell(2, col, _day_label(d, lang))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _header_fill()
        c.alignment = Alignment(horizontal="center")
        col += 2
    ws.cell(2, col, _tr(lang, "total_hours")).font = Font(bold=True, color="FFFFFF")
    ws.cell(2, col).fill = _header_fill()

    # Row 3 — IN/OUT sub-headers
    ws.cell(3, 1).fill = _header_fill()
    ws.cell(3, 2).fill = _header_fill()
    col = 3
    for _ in week_days:
        for label in [_tr(lang, "in"), _tr(lang, "out")]:
            c = ws.cell(3, col, label)
            c.font = Font(bold=True, color="FFFFFF", size=9)
            c.fill = _header_fill()
            c.alignment = Alignment(horizontal="center")
            col += 1
    ws.cell(3, col).fill = _header_fill()

    # Data rows
    row = 4
    for emp_id in sorted(data, key=lambda x: int(x) if x.isdigit() else x):
        rec = data[emp_id]
        ws.cell(row, 1, emp_id)
        ws.cell(row, 2, rec["name"])
        fill = _alt_fill() if row % 2 == 0 else None
        col = 3
        for d in week_days:
            day_rec = rec["days"].get(d.isoformat(), {})
            c_in  = ws.cell(row, col,   day_rec.get("in", ""))
            c_out = ws.cell(row, col+1, day_rec.get("out", ""))
            if fill:
                c_in.fill = fill
                c_out.fill = fill
            col += 2
        c_tot = ws.cell(row, col, rec["total_hours"])
        c_tot.font = Font(bold=True)
        if fill:
            c_tot.fill = fill
        row += 1

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 28
    for i in range(3, col+2):
        ws.column_dimensions[get_column_letter(i)].width = 8

    ws.freeze_panes = "C4"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── Department report ─────────────────────────────────────────────────────────

def export_department_report(start_date: date, end_date: date,
                              lang: str = "en",
                              dept_filter: str = None) -> io.BytesIO:
    """Department-level attendance summary, one sheet per department."""
    conn = get_conn()

    week_days = []
    d = start_date
    while d <= end_date:
        week_days.append(d)
        d += timedelta(days=1)

    dept_where = f"AND u.department_id = {int(dept_filter)}" if dept_filter and dept_filter.isdigit() else ""
    dept_rows = conn.execute(f"""
        SELECT COALESCE(d.name, '{_tr(lang, "no_dept")}') as dept,
               u.employee_id, u.name as emp_name
        FROM users u
        LEFT JOIN departments d ON d.id = u.department_id
        WHERE u.is_active = 1 {dept_where}
        ORDER BY dept, u.name
    """).fetchall()

    events = conn.execute("""
        SELECT employee_id, DATE(timestamp) as day,
               MIN(timestamp) as first_in, MAX(timestamp) as last_out
        FROM events
        WHERE DATE(timestamp) BETWEEN ? AND ?
          AND employee_id IS NOT NULL
        GROUP BY employee_id, day
    """, (start_date.isoformat(), end_date.isoformat())).fetchall()
    conn.close()

    by_emp_day = {}
    for e in events:
        by_emp_day.setdefault(e["employee_id"], {})[e["day"]] = {
            "in":  e["first_in"][11:16]  if e["first_in"]  else "",
            "out": e["last_out"][11:16] if e["last_out"] else "",
        }

    wb = Workbook()
    wb.remove(wb.active)

    depts = {}
    for r in dept_rows:
        depts.setdefault(r["dept"], []).append(r)

    for dept_name, employees in depts.items():
        ws = wb.create_sheet(title=dept_name[:31])

        total_cols = 2 + len(week_days)*2 + 1
        ws.merge_cells(start_row=1, start_column=1,
                       end_row=1, end_column=total_cols)
        c = ws.cell(1, 1, f"{dept_name} — {start_date.strftime('%d/%m/%Y')} "
                          f"{_tr(lang,'to')} {end_date.strftime('%d/%m/%Y')}")
        c.font = Font(bold=True, size=12, color="FFFFFF")
        c.fill = _header_fill()
        c.alignment = Alignment(horizontal="center")

        ws.cell(2, 1, "ID").font = Font(bold=True, color="FFFFFF")
        ws.cell(2, 1).fill = _header_fill()
        ws.cell(2, 2, _tr(lang, "name")).font = Font(bold=True, color="FFFFFF")
        ws.cell(2, 2).fill = _header_fill()
        col = 3
        for d in week_days:
            ws.merge_cells(start_row=2, start_column=col,
                           end_row=2, end_column=col+1)
            c = ws.cell(2, col, _day_label(d, lang))
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = _header_fill()
            c.alignment = Alignment(horizontal="center")
            col += 2
        ws.cell(2, col, _tr(lang, "days")).font = Font(bold=True, color="FFFFFF")
        ws.cell(2, col).fill = _header_fill()

        ws.cell(3, 1).fill = _header_fill()
        ws.cell(3, 2).fill = _header_fill()
        col = 3
        for _ in week_days:
            for label in [_tr(lang, "in"), _tr(lang, "out")]:
                c = ws.cell(3, col, label)
                c.font = Font(bold=True, color="FFFFFF", size=9)
                c.fill = _header_fill()
                col += 1
        ws.cell(3, col, _tr(lang, "days")).font = Font(bold=True, color="FFFFFF")
        ws.cell(3, col).fill = _header_fill()

        row = 4
        for emp in employees:
            ws.cell(row, 1, emp["employee_id"])
            ws.cell(row, 2, emp["emp_name"])
            fill = _alt_fill() if row % 2 == 0 else None
            days_present = 0
            col = 3
            emp_days = by_emp_day.get(str(emp["employee_id"]), {})
            for d in week_days:
                rec = emp_days.get(d.isoformat(), {})
                ws.cell(row, col,   rec.get("in", ""))
                ws.cell(row, col+1, rec.get("out", ""))
                if rec:
                    days_present += 1
                if fill:
                    ws.cell(row, col).fill = fill
                    ws.cell(row, col+1).fill = fill
                col += 2
            ws.cell(row, col, days_present)
            row += 1

        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 26
        for i in range(3, col+2):
            ws.column_dimensions[get_column_letter(i)].width = 7

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── Monthly summary ───────────────────────────────────────────────────────────

def export_monthly_summary(year: int, month: int,
                            lang: str = "en",
                            dept_filter: str = None) -> io.BytesIO:
    """Monthly attendance summary per employee."""
    conn = get_conn()
    days_in_month = calendar.monthrange(year, month)[1]
    start = date(year, month, 1)
    end   = date(year, month, days_in_month)

    no_dept = _tr(lang, "no_dept")
    dept_where = f"AND u.department_id = {int(dept_filter)}" if dept_filter and str(dept_filter).isdigit() else ""
    rows = conn.execute(f"""
        SELECT u.employee_id,
               COALESCE(d.name, '{no_dept}') as dept,
               u.name,
               COUNT(DISTINCT DATE(e.timestamp)) as days_present,
               SUM(CASE WHEN e.direction='IN'  THEN 1 ELSE 0 END) as total_in,
               SUM(CASE WHEN e.direction='OUT' THEN 1 ELSE 0 END) as total_out
        FROM users u
        LEFT JOIN departments d ON d.id = u.department_id
        LEFT JOIN events e ON e.employee_id = u.employee_id
            AND DATE(e.timestamp) BETWEEN ? AND ?
        WHERE u.is_active = 1 {dept_where}
        GROUP BY u.employee_id
        ORDER BY dept, u.name
    """, (start.isoformat(), end.isoformat())).fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d}"

    title = f"{_tr(lang, 'monthly_title')} — {_month_label(year, month, lang)}"
    ws.merge_cells("A1:H1")
    c = ws.cell(1, 1, title)
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = _header_fill()
    c.alignment = Alignment(horizontal="center")

    headers = [
        _tr(lang, "employee_id"),
        _tr(lang, "department"),
        _tr(lang, "name"),
        _tr(lang, "days_present"),
        f"/ {days_in_month}",
        _tr(lang, "attendance_pct"),
        _tr(lang, "in_count"),
        _tr(lang, "out_count"),
    ]
    _apply_header(ws, 2, headers)

    for i, r in enumerate(rows, 3):
        pct = round((r["days_present"] or 0) / days_in_month * 100, 1)
        fill = _alt_fill() if i % 2 == 0 else None
        for col, val in enumerate([
            r["employee_id"], r["dept"], r["name"],
            r["days_present"] or 0, days_in_month, f"{pct}%",
            r["total_in"] or 0, r["total_out"] or 0,
        ], 1):
            c = ws.cell(i, col, val)
            c.border = _thin_border()
            c.alignment = Alignment(horizontal="center")
            if fill:
                c.fill = fill

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 28
    for col in "DEFGH":
        ws.column_dimensions[col].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
